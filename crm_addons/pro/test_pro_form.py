"""The lead form follows CRM's layouts; import, export and the compatibility check."""

import json

import frappe

from crm_addons.pro import compat, form, leads
from crm_addons.team_access.test_access import AccessCase


class TestProForm(AccessCase):
	def test_create_form_is_crm_quick_entry(self):
		frappe.set_user("Administrator")
		names = {f["fieldname"] for s in form.get_form("create") for f in s["fields"]}
		quick = {n for _l, ns in form._layout_fieldnames("Quick Entry") for n in ns}
		self.assertTrue(names)
		self.assertTrue(names <= quick | {df.fieldname for df in frappe.get_meta("CRM Lead").fields if df.reqd})

	def test_values_outside_the_layouts_are_refused(self):
		frappe.set_user("Administrator")
		for bad in ("converted", "owner", "sla_status", "naming_series"):
			with self.assertRaises(frappe.ValidationError):
				form.clean_values({bad: "1"})
		self.assertEqual(form.clean_values({"first_name": " A ", "email": ""}), {"first_name": "A", "email": None})

	def test_export_has_every_field_with_labels(self):
		frappe.set_user("Administrator")
		cols = dict(form.export_fields())
		for k in ("name", "lead_name", "first_name", "mobile_no", "status", "lead_owner", "organization", "creation"):
			self.assertIn(k, cols)
		self.assertNotIn("status_change_log", cols)

	def test_import_dry_run_then_real(self):
		frappe.set_user("Administrator")
		tag = frappe.generate_hash(length=6)
		text = f"First Name,Last Name,Email,Nonsense\nZq,{tag},zq{tag}@example.com,x\n,missing first,,\n"
		dry = form.import_csv(text, 1)
		self.assertEqual((dry["created"], len(dry["failed"]), dry["unknown"]), (1, 1, ["Nonsense"]))
		self.assertFalse(frappe.db.exists("CRM Lead", {"last_name": tag}))
		real = form.import_csv(text, 0)
		name = frappe.db.get_value("CRM Lead", {"last_name": tag})
		self.assertTrue(name)
		self.addCleanup(frappe.delete_doc, "CRM Lead", name, force=True, ignore_permissions=True)
		self.assertEqual(real["created"], 1)
		with self.assertRaises(frappe.ValidationError):
			form.import_csv("Nothing,Matches\n1,2", 1)

	def test_template_headers_import_cleanly(self):
		frappe.set_user("Administrator")
		self.assertTrue(form.import_template()["columns"])

	def test_selected_export_and_guest(self):
		frappe.set_user("Administrator")
		doc = frappe.get_doc({"doctype": "CRM Lead", "first_name": "Zq", "last_name": "Pick"}).insert(ignore_permissions=True)
		self.addCleanup(frappe.delete_doc, "CRM Lead", doc.name, force=True, ignore_permissions=True)
		res = leads.export_csv(names=json.dumps([doc.name]))
		self.assertEqual(res["count"], 1)
		frappe.set_user("Guest")
		with self.assertRaises(frappe.PermissionError):
			form.get_form()

	def test_compat_is_clean_here_and_notices_changes(self):
		frappe.set_user("Administrator")
		self.assertEqual(compat.problems(), [])
		now = compat._snapshot()
		old = json.loads(json.dumps(now))
		old["lead_fields"]["gone_field"] = "Data"
		del old["lead_fields"]["status"]
		old["versions"]["crm"] = "0.0.1"
		text = " ".join(compat.changes(old, now))
		self.assertIn("gone_field", text)
		self.assertIn("status", text)
		self.assertIn("0.0.1", text)
		self.assertEqual(compat.changes(None, now), [])

	def test_compat_reports_a_missing_field(self):
		frappe.set_user("Administrator")
		compat.NEEDS["CRM Lead"].append("zz_not_a_field")
		self.addCleanup(compat.NEEDS["CRM Lead"].remove, "zz_not_a_field")
		self.assertTrue(any("zz_not_a_field" in p["text"] and p["level"] == "error" for p in compat.problems()))

	def test_export_columns_for_the_download_button(self):
		frappe.set_user("Administrator")
		cols = form.export_columns()
		self.assertEqual(cols[0], {"key": "name", "label": "ID"})
		self.assertTrue(any(c["key"] == "mobile_no" for c in cols))


class TestToXlsx(AccessCase):
	def test_builds_a_workbook_and_quotes_formulas(self):
		import base64
		import io

		import openpyxl

		frappe.set_user("Administrator")
		raw = form.to_xlsx(json.dumps(["Name", "Phone"]), json.dumps([["=cmd()", "+919800000000"], ["Ann", None]]))
		ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(raw))).active
		self.assertEqual([c.value for c in ws[1]], ["Name", "Phone"])
		self.assertEqual(ws["A2"].value, "'=cmd()")
		self.assertEqual(ws["B2"].value, "+919800000000")
		frappe.set_user("Guest")
		with self.assertRaises(frappe.PermissionError):
			form.to_xlsx("[]", "[]")
